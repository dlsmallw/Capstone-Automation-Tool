import os

import pandas as pd
import openpyxl as opyxl

import tkinter as tk
from tkinter import filedialog

root = tk.Tk()
root.withdraw()

def get_us_df(us_fp):
    raw_df = None

    if us_fp.split(".")[-1] == "csv":
        raw_df = pd.read_csv(us_fp)
    elif us_fp.split(".")[-1] == "xlsx":
        raw_df = pd.read_csv(us_fp)
    else:
        return None
    
    if raw_df is not None:
        raw_df = raw_df[['ref', 'sprint', 'total-points']]

    return raw_df

def get_task_df(task_fp):
    raw_df = None

    if task_fp.split(".")[-1] == "csv":
        raw_df = pd.read_csv(task_fp)
        
    elif task_fp.split(".")[-1] == "xlsx":
        raw_df = pd.read_csv(task_fp)
        return raw_df[['ref', 'subject', 'user_story', 'sprint', 'assigned_to']]
    
    if raw_df is not None:
        raw_df = raw_df[['ref', 'subject', 'user_story', 'sprint', 'assigned_to']]
        raw_df.sort_values(['sprint'], ascending=[True], inplace=True)
    
    return raw_df

    
def file_is_valid(filename):
    if filename != "" and filename is not None:
        return True
    else:
        return False
    
def make_hyperlink(value):
    base_url = 'https://tree.taiga.io/project/dlsmallw-group-8-asu-capstone-natural-language-processing-for-decolonizing-harm-reduction-literature/task/{}'
    return '=HYPERLINK("%s", "Task-%s")' % (base_url.format(value), value)

def get_members(dataframe):
    members = []

    for row in dataframe:
        if pd.notnull(row):
            members.append(row) if row not in members else None

    return members

def get_sprints(dataframe):
    sprints = []

    for row in dataframe:
        if pd.notnull(row):
            sprints.append(row) if row not in sprints else None

    return sprints

def create_new_wb(filename, sprints):
    if (os.path.exists(filename)):
        os.remove(filename)
    
    wb = opyxl.Workbook()

    wb.create_sheet("All_Data")

    for sprint in sprints:
        wb.create_sheet(sprint)

    for sheet in wb.sheetnames:
        if sheet not in sprints and sheet != "All_Data":
            del wb[sheet]

    wb.save(filename)

def write_data(df, writer, sheet):
    df.to_excel(writer, sheet_name=sheet)

def format_df(df):
    for i, row in df.iterrows():
        us_num = row['user_story']
        task_num = row['task']

        df.at[i, 'user_story'] = f'US-{int(us_num)}' if pd.notnull(us_num) else "Storyless"
        df.at[i, 'task'] = make_hyperlink(task_num)

    return df

def parse_data(us_fp, task_fp):
    filename = 'taiga_parse_data.xlsx'
    data_columns = ['subject', 'sprint', 'user_story', 'points', 'task', 'coding']

    us_df = get_us_df(us_fp)
    task_df = get_task_df(task_fp)

    members = get_members(task_df['assigned_to'])
    num_mems = len(members)

    data_columns.extend(members)

    sprints = get_sprints(us_df['sprint'])
    create_new_wb(filename, sprints)

    all_data = [0] * len(task_df)

    for index, row in task_df.iterrows():
        assigned = row['assigned_to']

        subject = row['subject']
        sprint = row['sprint']

        us_num = row['user_story']
        us_row = us_df.loc[us_df['ref'] == us_num]

        # user_story = f'US-{int(us_num)}' if pd.notnull(us_num) else "Storyless"
        user_story = row['user_story']
        points = us_row['total-points'].values[0] if pd.notnull(us_num) else 0

        # task = make_hyperlink(row['ref'])
        task = row['ref']
        coding = ""

        mem_data = [None] * num_mems

        i = 0
        for mem in members:
            mem_data[i] = "100%" if assigned == mem else None
            i += 1

        data_row = [subject, sprint, user_story, points, task, coding]
        data_row.extend(mem_data)

        all_data[index] = data_row

    parsed_df = pd.DataFrame(all_data, columns=data_columns)

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        write_data(format_df(parsed_df.sort_values(['sprint', 'user_story', 'task'], ascending=[True, True, True])), writer, "All_Data")
        
        for sprint in sprints:
            sprint_df = parsed_df.loc[parsed_df['sprint'] == sprint]
            write_data(format_df(sprint_df.sort_values(['sprint', 'user_story', 'task'], ascending=[True, True, True])), writer, sprint)
        
us_taiga_fp = filedialog.askopenfilename()
task_taiga_fp = filedialog.askopenfilename()

if file_is_valid(us_taiga_fp) and file_is_valid(task_taiga_fp):
    parse_data(us_taiga_fp, task_taiga_fp)
else:
    print("Invalid files selected")